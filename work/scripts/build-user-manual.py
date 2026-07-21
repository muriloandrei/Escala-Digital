from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
ASSETS = DOCS / 'manual-usuario-assets'
OUT = DOCS / 'manual-usuario-escala-inteligente.docx'
LOGO = ROOT / 'public' / 'assets' / 'escala-inteligente-logo.png'

NAVY = '0B2545'
BLUE = '2563EB'
MUTED = '64748B'
LIGHT = 'E8EEF5'
BORDER = 'CBD5E1'

screens = {
    'login': ASSETS / '01-login.png',
    'escalas': ASSETS / '02-escalas-geradas.png',
    'funcionarios': ASSETS / '03-funcionarios.png',
    'secoes': ASSETS / '04-secoes.png',
    'secao_form': ASSETS / '05-secao-formulario.png',
    'turnos': ASSETS / '06-turnos-secao.png',
    'turno_form': ASSETS / '07-turno-secao-formulario.png',
    'esc_func': ASSETS / '08-escalas-funcionarios.png',
    'historico': ASSETS / '09-historico.png',
    'tipos': ASSETS / '10-tipos-descanso.png',
    'acessos': ASSETS / '11-controle-acesso.png',
    'perfis': ASSETS / '12-perfis-acesso.png',
    'config': ASSETS / '13-configuracoes.png',
    'modal_criar': ASSETS / '14-modal-criar-escala.png',
    'abrir': ASSETS / '15-abrir-escala.png',
    'criacao_secoes': ASSETS / '16-criacao-selecionar-secoes.png',
    'criacao_timeline': ASSETS / '17-criacao-timeline-gerada.png',
    'criacao_distribuir': ASSETS / '18-criacao-distribuir-funcionarios.png',
    'criacao_folgas': ASSETS / '19-criacao-folgas-distribuidas.png',
    'criacao_detalhada': ASSETS / '20-criacao-escala-detalhada.png',
}

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def set_cell_border(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)

def style_table(table, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.name = 'Calibri'
                    run.font.size = Pt(9)
            if header and row_idx == 0:
                set_cell_shading(cell, LIGHT)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor.from_string(NAVY)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f'Heading {level}'
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE if level < 3 else NAVY)
    return p

def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(NAVY)
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)

def add_steps(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)

def add_note(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0,0)
    set_cell_shading(cell, 'F8FAFC')
    set_cell_border(cell, 'BFDBFE')
    p = cell.paragraphs[0]
    r = p.add_run(title + ': ')
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p.add_run(text)
    doc.add_paragraph()

def add_screenshot(doc, key, caption):
    image = screens.get(key)
    if not image or not image.exists():
        add_note(doc, 'Print pendente', f'Imagem nao encontrada para: {caption}')
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image), width=Inches(7.0))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    for run in cap.runs:
        run.font.size = Pt(9)
        run.font.italic = True
        run.font.color.rgb = RGBColor.from_string(MUTED)

def add_button_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    hdr = table.rows[0].cells
    hdr[0].text = 'Botao ou acao'
    hdr[1].text = 'Onde fica'
    hdr[2].text = 'O que faz'
    for action, where, desc in rows:
        cells = table.add_row().cells
        cells[0].text = action
        cells[1].text = where
        cells[2].text = desc
    style_table(table)
    doc.add_paragraph()

def build():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.6)
    sec.right_margin = Inches(0.6)
    sec.header_distance = Inches(0.3)
    sec.footer_distance = Inches(0.3)

    styles = doc.styles
    styles['Normal'].font.name = 'Calibri'
    styles['Normal'].font.size = Pt(10.5)
    styles['Normal'].paragraph_format.space_after = Pt(6)
    for style_name, size, color in [('Heading 1',16,BLUE),('Heading 2',13,BLUE),('Heading 3',12,NAVY)]:
        st = styles[style_name]
        st.font.name = 'Calibri'
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(5)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run('Escala Inteligente Savegnago | Manual do Usuario')
    hr.font.size = Pt(8)
    hr.font.color.rgb = RGBColor.from_string(MUTED)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run('Equipe de Desenvolvimento Grupo Savegnago')
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor.from_string(MUTED)

    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Inches(2.1))
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run('Manual de Uso - Escala Inteligente Savegnago')
    tr.font.name = 'Calibri'
    tr.font.size = Pt(24)
    tr.font.bold = True
    tr.font.color.rgb = RGBColor.from_string(NAVY)
    sub = doc.add_paragraph('Guia operacional para login, criação, acompanhamento, edição e configuração das escalas.')
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    add_note(doc, 'Observacao', 'As telas e permissoes podem variar conforme o perfil de acesso e as lojas vinculadas ao usuario. Nao informe senhas em chamados ou documentos; use sempre credenciais fornecidas pelo administrador.')

    add_heading(doc, '1. Acessar o sistema')
    add_screenshot(doc, 'login', 'Tela de login da aplicacao.')
    add_steps(doc, [
        'Abra o endereco da aplicacao no navegador.',
        'Digite seu usuario no campo Login.',
        'Digite sua senha no campo Senha.',
        'Clique em Escala Inteligente para entrar.',
        'Ao finalizar o uso, clique em Sair no canto superior direito.'
    ])

    add_heading(doc, '2. Navegacao principal')
    add_body(doc, 'A aplicacao usa uma estrutura de dashboard: menu lateral para modulos, breadcrumb no topo para indicar a tela atual e area principal para filtros, tabelas e formularios.')
    add_bullets(doc, [
        'Escalas: concentra Escalas Geradas, Escalas por Funcionario, Secoes e Turnos por Secao.',
        'Funcionarios: consulta funcionarios cadastrados para as lojas permitidas.',
        'Configuracoes: controle de acesso, perfis e regras gerais da aplicacao.',
        'Badge de loja: mostra a quantidade de lojas vinculadas ao usuario.',
        'Avatar do usuario: identifica o usuario logado; o botao Sair encerra a sessao.'
    ])

    add_heading(doc, '3. Escalas Geradas')
    add_screenshot(doc, 'escalas', 'Tela Escalas Geradas, com pesquisa, filtros e acoes por escala.')
    add_body(doc, 'Use esta tela como ponto central da operacao. Ao entrar nela, o sistema consulta o banco e lista as escalas mais recentes conforme as lojas do usuario.')
    add_button_table(doc, [
        ('Buscar mes ou loja', 'Barra de filtros', 'Filtra a listagem digitando parte do mes ou da loja.'),
        ('Filtro Loja', 'Barra de filtros', 'Mostra todas as lojas permitidas ou apenas uma loja especifica.'),
        ('Filtro Mes/Ano', 'Barra de filtros', 'Restringe a consulta ao periodo selecionado.'),
        ('Filtro Status', 'Barra de filtros', 'Filtra escalas Ativas, Agendadas, Modificadas ou Finalizadas.'),
        ('Criar Nova Escala', 'Cabecalho da pagina', 'Abre o inicio do fluxo de criacao de uma nova escala.'),
        ('Abrir Escala', 'Coluna Acoes', 'Abre a escala existente para visualizacao e edicao conforme permissao.'),
        ('Oficializar', 'Coluna Acoes', 'Marca a revisao atual como pronta. Novas alteracoes exigem nova oficializacao.'),
        ('Historico', 'Coluna Acoes', 'Abre os registros de revisoes e alteracoes da escala.'),
        ('Excluir/Inativar', 'Coluna Acoes', 'Inativa a escala sem apagar fisicamente os registros do banco.')
    ])

    add_heading(doc, '4. Criar uma nova escala')
    add_screenshot(doc, 'modal_criar', 'Print 14 - Modal inicial para escolher loja, mes e ano da escala.')
    add_steps(doc, [
        'Na tela Escalas Geradas, clique em Criar Nova Escala.',
        'Escolha a loja, o mes e o ano.',
        'Clique em Iniciar Criacao.',
        'Se ja existir escala ativa para a mesma loja e mes, o sistema bloqueia a duplicidade e oferece Abrir Escala.',
        'Se nao existir, o sistema abre a pagina de criacao com secoes e timeline.'
    ])
    add_note(doc, 'Regra importante', 'Uma loja deve ter apenas uma escala ativa por mes. Escalas inativas sao ignoradas para essa validacao.')

    add_heading(doc, '5. Fluxo completo do Iniciar Criacao')
    add_body(doc, 'Depois de clicar em Iniciar Criacao, o sistema abre a pagina de rascunho da escala. E aqui que voce escolhe as secoes, gera a timeline, distribui os funcionarios/folgas e confirma a escala detalhada antes de salvar no banco.')
    add_screenshot(doc, 'criacao_secoes', 'Print 16 - Etapa 1: selecionar as secoes/turnos que farao parte da escala.')
    add_steps(doc, [
        'Confira loja e periodo no status do rascunho.',
        'Marque as secoes/turnos que devem compor a escala.',
        'Deixe desmarcado qualquer turno que nao deve entrar naquele mes.',
        'Clique em Atualizar Timeline para montar a linha do tempo.'
    ])
    add_screenshot(doc, 'criacao_timeline', 'Print 17 - Etapa 2: timeline gerada com as secoes selecionadas.')
    add_body(doc, 'A timeline mostra os blocos de trabalho, intervalo e cobertura por secao. Antes de distribuir funcionarios, revise se os horarios e quantidades estao corretos.')
    add_screenshot(doc, 'criacao_distribuir', 'Print 18 - Etapa 3: tela de distribuicao dos funcionarios antes da regra 5x2.')
    add_body(doc, 'Ao clicar em Distribuir Funcionarios, o sistema abre a grade operacional da escala. Nessa tela voce confere colaboradores, dias do mes e a barra de botoes da escala.')
    add_screenshot(doc, 'criacao_folgas', 'Print 19 - Etapa 4: folgas distribuidas pela regra 5x2 e timeline bloqueada para edicao.')
    add_body(doc, 'Depois de clicar em Distribuir Folgas 5x2, a timeline fica bloqueada para impedir mudancas sem nova validacao. O botao Gerar Escala Detalhada passa a aparecer. Para mudar colaboradores ou folgas, clique em Editar e distribua novamente antes de gerar a detalhada.')
    add_screenshot(doc, 'criacao_detalhada', 'Print 20 - Etapa 5: escala detalhada pronta para validar, salvar ou imprimir.')
    add_button_table(doc, [
        ('Atualizar Timeline', 'Pagina Nova Escala', 'Monta ou remonta a timeline com as secoes/turnos marcados.'),
        ('Remover secao', 'Linha da timeline', 'Remove uma secao/turno do rascunho antes da distribuicao.'),
        ('Imprimir Timeline', 'Cabecalho da timeline', 'Imprime a visualizacao da timeline de turnos.'),
        ('Distribuir Funcionarios', 'Cabecalho da timeline', 'Abre a grade/esqueleto para distribuir funcionarios, folgas e gerar a escala.'),
        ('Carregar Funcionarios', 'Modal Escala', 'Carrega os funcionarios da loja quando a grade precisar ser atualizada.'),
        ('Distribuir Folgas 5x2', 'Modal Escala', 'Aplica automaticamente as folgas conforme a regra 5x2 e bloqueia a edicao da timeline.'),
        ('Editar', 'Modal Escala apos distribuicao', 'Libera a edicao novamente, esconde Gerar Escala Detalhada e exige nova distribuicao.'),
        ('Visualizar Timeline', 'Modal Escala', 'Abre uma visualizacao ampliada da timeline usada na escala.'),
        ('Gerar Escala Detalhada', 'Modal Escala apos distribuicao', 'Converte a grade distribuida em escala diaria por funcionario.'),
        ('Imprimir Esqueleto', 'Modal Escala', 'Imprime a grade/esqueleto de distribuicao.'),
        ('Validar Escala', 'Modal Escala Detalhada', 'Executa as regras de validacao antes do salvamento.'),
        ('Salvar Escala', 'Modal Escala Detalhada', 'Grava a escala no banco e retorna para Escalas Geradas.'),
        ('Imprimir Escala Detalhada', 'Modal Escala Detalhada', 'Imprime a escala mensal detalhada.'),
        ('Fechar', 'Modais de escala', 'Fecha a janela atual. Se houver rascunho nao salvo, o progresso pode ser perdido ao sair da tela.')
    ])

    add_heading(doc, '5.1 Sequencia visual dos prints 15 a 20')
    add_body(doc, 'Esta sequencia reune os prints especificos do fluxo de criacao e abertura da escala para consulta rapida.')
    add_screenshot(doc, 'abrir', 'Print 15 - Abrir uma escala existente para visualizacao/edicao.')
    add_screenshot(doc, 'criacao_secoes', 'Print 16 - Selecionar secoes/turnos no rascunho da nova escala.')
    add_screenshot(doc, 'criacao_timeline', 'Print 17 - Gerar e revisar a timeline da escala.')
    add_screenshot(doc, 'criacao_distribuir', 'Print 18 - Distribuir funcionarios na grade/esqueleto.')
    add_screenshot(doc, 'criacao_folgas', 'Print 19 - Aplicar Distribuir Folgas 5x2 e bloquear edicao.')
    add_screenshot(doc, 'criacao_detalhada', 'Print 20 - Gerar a escala detalhada para validar, salvar e imprimir.')

    add_heading(doc, '6. Abrir e editar uma escala pronta')
    add_screenshot(doc, 'abrir', 'Print 15 - Visualizacao ao abrir uma escala existente.')
    add_body(doc, 'Ao abrir uma escala, o sistema apresenta as secoes em abas. Cada aba concentra a timeline e a escala detalhada daquela secao, facilitando a conferencia e edicao por grupo de colaboradores.')
    add_bullets(doc, [
        'Use as abas para alternar entre secoes da escala.',
        'Edite dias/horarios apenas quando a escala ainda permitir modificacao.',
        'Toda alteracao relevante cria uma nova revisao e exige nova oficializacao.',
        'Escalas finalizadas ou inativadas nao devem ser editadas.'
    ])

    add_heading(doc, '7. Escalas por Funcionario')
    add_screenshot(doc, 'esc_func', 'Tela Escalas por Funcionario, usada para localizar e editar a escala individual.')
    add_body(doc, 'Use esta tela quando a manutencao for individual, por funcionario. Os filtros permitem localizar funcionarios por loja, mes e outros criterios disponiveis.')
    add_button_table(doc, [
        ('Pesquisar', 'Filtros da pagina', 'Localiza funcionario por nome, chapa, secao ou funcao.'),
        ('Filtro Loja/Mes/Ano', 'Filtros da pagina', 'Define o periodo e lojas exibidos.'),
        ('Editar', 'Coluna Acoes', 'Abre a escala detalhada do funcionario selecionado.'),
        ('Salvar', 'Tela de edicao do funcionario', 'Grava alteracoes e cria nova revisao quando necessario.'),
        ('Imprimir Escala', 'Tela de edicao do funcionario', 'Imprime a escala mensal individual com campo de ciencia/assinatura.')
    ])
    add_body(doc, 'Na edicao por funcionario, clique no dia da grade para alterar o tipo do dia. Se escolher Descanso, selecione o tipo de descanso; se escolher Trabalho, selecione um turno cadastrado ou informe horarios manualmente. A justificativa da mudanca deve ser preenchida.')

    add_heading(doc, '8. Funcionarios')
    add_screenshot(doc, 'funcionarios', 'Tela Funcionarios, com filtros e coluna de acoes.')
    add_bullets(doc, [
        'Mostra os funcionarios cadastrados no banco para as lojas permitidas.',
        'Permite pesquisar por nome, chapa, secao ou funcao.',
        'Os campos Secao e Funcao exibem as descricoes para facilitar a leitura.',
        'O botao Editar abre a manutencao do cadastro conforme permissao do perfil.'
    ])

    add_heading(doc, '9. Secoes')
    add_screenshot(doc, 'secoes', 'Lista de secoes cadastradas por loja.')
    add_screenshot(doc, 'secao_form', 'Formulario para criar ou editar secao.')
    add_steps(doc, [
        'Acesse Escalas > Secoes.',
        'Use a pesquisa para localizar por codigo ou descricao.',
        'Filtre por Todas as Lojas ou por uma loja especifica.',
        'Clique em Nova Secao para cadastrar ou Editar para alterar uma secao existente.',
        'Salve para gravar os dados em SGN_ESC_SECAO.'
    ])

    add_heading(doc, '10. Turnos por Secao')
    add_screenshot(doc, 'turnos', 'Lista de turnos cadastrados para cada secao.')
    add_screenshot(doc, 'turno_form', 'Formulario de turno por secao.')
    add_body(doc, 'Os turnos por secao alimentam o fluxo de criacao da escala. Cada secao pode ter varios turnos, com quantidade prevista de colaboradores e horarios de entrada/saida.')
    add_button_table(doc, [
        ('Novo Turno', 'Cabecalho da pagina', 'Abre o formulario de cadastro de turno por secao.'),
        ('Filtro Loja', 'Barra de filtros', 'Filtra turnos pelas lojas permitidas.'),
        ('Filtro Secao', 'Barra de filtros', 'Mostra todos os turnos ou uma secao especifica.'),
        ('Editar', 'Coluna Acoes', 'Altera horarios ou quantidade prevista do turno.')
    ])

    add_heading(doc, '11. Tipos de Descanso')
    add_screenshot(doc, 'tipos', 'Tela Tipos de Descanso.')
    add_body(doc, 'Cadastre os tipos usados na edicao de dias de descanso, como Folga ou Ferias. Cada tipo possui descricao, sigla e status. A sigla aparece nas grades de escala.')

    add_heading(doc, '12. Historico')
    add_screenshot(doc, 'historico', 'Relatorio de historico e revisoes.')
    add_body(doc, 'A pagina de Historico registra alteracoes passo a passo: criacao, edicao de horarios, inclusao/remocao de funcionario, revisoes, oficializacoes e usuario responsavel. Use os filtros para auditar uma loja ou periodo especifico.')

    add_heading(doc, '13. Controle de Acesso')
    add_screenshot(doc, 'acessos', 'Tela Controle de Acesso.')
    add_body(doc, 'Nesta tela, administradores gerenciam usuarios, lojas vinculadas e perfil de acesso. O vinculo de lojas limita quais lojas aparecem nos filtros e operacoes do usuario.')
    add_bullets(doc, [
        'Criar usuario: cadastra login, nome, senha inicial, perfil e lojas permitidas.',
        'Editar usuario: altera perfil, status e lojas vinculadas.',
        'Inativar usuario: remove o acesso sem excluir o historico.'
    ])

    add_heading(doc, '14. Perfil de Acesso')
    add_screenshot(doc, 'perfis', 'Tela Perfil de Acesso.')
    add_body(doc, 'Perfis controlam o que cada grupo pode visualizar, editar e inativar em cada pagina. Use perfis para separar usuarios administrativos, operadores de loja e consulta.')

    add_heading(doc, '15. Configuracoes')
    add_screenshot(doc, 'config', 'Tela Configuracoes.')
    add_body(doc, 'Concentra regras gerais da aplicacao, parametros da linha do tempo e regras de turno. Alteracoes aqui impactam os calculos e validacoes, por isso devem ser feitas apenas por usuarios autorizados.')

    add_heading(doc, '16. Boas praticas operacionais')
    add_bullets(doc, [
        'Cadastre ou revise Secoes e Turnos por Secao antes de criar escalas.',
        'Sempre valide a escala antes de salvar ou oficializar.',
        'Use Oficializar somente quando a escala estiver pronta para uso.',
        'Nao delete dados diretamente no banco; use inativacao pelo sistema quando disponivel.',
        'Ao alterar uma escala ja oficializada, confira a nova revisao e oficialize novamente.',
        'Use Historico para rastrear quem alterou, quando alterou e o que mudou.'
    ])

    add_heading(doc, '17. Resumo das principais tabelas')
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = 'Tabela'
    table.rows[0].cells[1].text = 'Uso principal na aplicacao'
    rows = [
        ('SGN_ESC_USUARIO / SGN_ESC_USUARIO_LOJA', 'Login, perfil e lojas permitidas do usuario.'),
        ('SGN_ESC_FUNCIONARIO', 'Cadastro base dos funcionarios consultados por loja.'),
        ('SGN_ESC_SECAO', 'Cadastro de secoes por loja.'),
        ('SGN_ESC_SECAO_TURNO', 'Turnos e quantidade prevista por secao.'),
        ('SGN_ESC_PROG', 'Cabecalho da programacao mensal por funcionario/revisao.'),
        ('SGN_ESC_PROG_DIA', 'Dias, horarios e programacao da escala.'),
        ('SGN_ESC_TIPO_DESCANSO', 'Tipos de descanso usados na edicao da escala.'),
        ('Tabelas de auditoria', 'Registro de alteracoes, revisoes e oficializacoes.')
    ]
    for a,b in rows:
        c = table.add_row().cells
        c[0].text = a
        c[1].text = b
    style_table(table)

    DOCS.mkdir(exist_ok=True)
    doc.save(OUT)
    print(OUT)

if __name__ == '__main__':
    build()


